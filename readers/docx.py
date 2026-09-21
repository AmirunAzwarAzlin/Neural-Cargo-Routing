import io

from docx import Document

from core.models import DocumentExtraction
from llm.gemini_client import extract_from_text
from readers.dispatch import register
from readers.zip_safety import UnsafeZipError, check_zip_bomb_safety

MAX_TEXT_CHARS = 100_000


def read_docx(data: bytes, role: str, filename: str) -> DocumentExtraction:
    try:
        check_zip_bomb_safety(data)
        doc = Document(io.BytesIO(data))
    except (UnsafeZipError, Exception):
        return DocumentExtraction(role=role, filename=filename, doc_kind="UNKNOWN", readable=False)

    lines = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            lines.append(" | ".join(cells))
    text = "\n".join(lines).strip()

    if not text:
        return DocumentExtraction(role=role, filename=filename, doc_kind="UNKNOWN", readable=False)

    if len(text) > MAX_TEXT_CHARS:
        # Can't confidently read the whole document — escalate rather than
        # silently sending a truncated/costly prompt to Gemini.
        return DocumentExtraction(role=role, filename=filename, doc_kind="UNKNOWN", readable=False)

    return extract_from_text(data, text, role, filename)


register("docx", read_docx)
