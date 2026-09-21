"""Reader-dispatch tests for non-.txt formats. Gemini calls are monkeypatched
so this suite runs offline, deterministically, and at no API cost."""
import io

import openpyxl
from docx import Document

from core.models import DocumentExtraction, FieldState


def test_docx_reader_extracts_paragraphs_and_tables_then_calls_gemini(monkeypatch):
    doc = Document()
    doc.add_paragraph("BILL OF LADING (DRAFT)")
    doc.add_paragraph("Shipper: TEST SHIPPER CO")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Consignee"
    table.rows[0].cells[1].text = "TEST CONSIGNEE CO"
    buf = io.BytesIO()
    doc.save(buf)
    data = buf.getvalue()

    captured = {}

    def fake_extract_from_text(document_bytes, document_text, role, filename):
        captured["text"] = document_text
        captured["role"] = role
        return DocumentExtraction(role=role, filename=filename, doc_kind="BL", readable=True)

    import readers.docx as docx_reader
    monkeypatch.setattr(docx_reader, "extract_from_text", fake_extract_from_text)

    result = docx_reader.read_docx(data, "BL", "att.docx")

    assert "Shipper: TEST SHIPPER CO" in captured["text"]
    assert "Consignee | TEST CONSIGNEE CO" in captured["text"]
    assert result.doc_kind == "BL"


def test_docx_reader_marks_unreadable_on_corrupt_bytes():
    import readers.docx as docx_reader

    result = docx_reader.read_docx(b"not a real docx", "SI", "att.docx")
    assert result.readable is False
    assert result.doc_kind == "UNKNOWN"


def test_xlsx_reader_dumps_cells_and_calls_gemini(monkeypatch):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "Shipper"
    ws["B1"] = "TEST SHIPPER CO"
    buf = io.BytesIO()
    wb.save(buf)
    data = buf.getvalue()

    captured = {}

    def fake_extract_from_text(document_bytes, document_text, role, filename):
        captured["text"] = document_text
        return DocumentExtraction(role=role, filename=filename, doc_kind="SI", readable=True)

    import readers.xlsx as xlsx_reader
    monkeypatch.setattr(xlsx_reader, "extract_from_text", fake_extract_from_text)

    result = xlsx_reader.read_xlsx(data, "SI", "att.xlsx")

    assert "Sheet1!A1: Shipper" in captured["text"]
    assert "Sheet1!B1: TEST SHIPPER CO" in captured["text"]
    assert result.doc_kind == "SI"


def test_xlsx_reader_marks_unreadable_on_corrupt_bytes():
    import readers.xlsx as xlsx_reader

    result = xlsx_reader.read_xlsx(b"not a real xlsx", "BL", "att.xlsx")
    assert result.readable is False


def test_xlsx_reader_skips_hidden_sheets(monkeypatch):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Visible"
    ws["A1"] = "Shipper"
    hidden = wb.create_sheet("Hidden")
    hidden.sheet_state = "hidden"
    hidden["A1"] = "SECRET DEFECT VALUE"
    buf = io.BytesIO()
    wb.save(buf)
    data = buf.getvalue()

    captured = {}

    def fake_extract_from_text(document_bytes, document_text, role, filename):
        captured["text"] = document_text
        return DocumentExtraction(role=role, filename=filename, doc_kind="SI", readable=True)

    import readers.xlsx as xlsx_reader
    monkeypatch.setattr(xlsx_reader, "extract_from_text", fake_extract_from_text)

    xlsx_reader.read_xlsx(data, "SI", "att.xlsx")

    assert "SECRET DEFECT VALUE" not in captured["text"]
    assert "Visible!A1: Shipper" in captured["text"]


def test_xlsx_reader_escalates_when_cell_count_exceeds_cap(monkeypatch):
    wb = openpyxl.Workbook()
    ws = wb.active
    for i in range(5):
        ws.cell(row=i + 1, column=1, value=f"v{i}")
    buf = io.BytesIO()
    wb.save(buf)
    data = buf.getvalue()

    import readers.xlsx as xlsx_reader
    monkeypatch.setattr(xlsx_reader, "MAX_CELLS", 2)

    result = xlsx_reader.read_xlsx(data, "SI", "att.xlsx")

    assert result.readable is False
    assert result.doc_kind == "UNKNOWN"


def test_docx_reader_rejects_zip_bomb(monkeypatch):
    import readers.docx as docx_reader
    from readers.zip_safety import UnsafeZipError

    def boom(data, **kwargs):
        raise UnsafeZipError("too big")

    monkeypatch.setattr(docx_reader, "check_zip_bomb_safety", boom)

    result = docx_reader.read_docx(b"PK\x03\x04fake", "SI", "att.docx")

    assert result.readable is False
    assert result.doc_kind == "UNKNOWN"


def test_pdf_reader_uses_text_path_when_text_present(monkeypatch):
    import readers.pdf as pdf_reader

    monkeypatch.setattr(pdf_reader, "_extract_text_per_page", lambda data: (["Shipper: TEST CO\n" * 5], False))
    called = {}

    def fake_extract_from_text(document_bytes, document_text, role, filename):
        called["used_text_path"] = True
        return DocumentExtraction(role=role, filename=filename, doc_kind="SI", readable=True)

    monkeypatch.setattr(pdf_reader, "extract_from_text", fake_extract_from_text)
    result = pdf_reader.read_pdf(b"fake-pdf-bytes", "SI", "att.pdf")

    assert called.get("used_text_path") is True
    assert result.doc_kind == "SI"


def test_pdf_reader_falls_back_to_vision_when_no_text(monkeypatch):
    import readers.pdf as pdf_reader

    monkeypatch.setattr(pdf_reader, "_extract_text_per_page", lambda data: ([""], False))
    monkeypatch.setattr(pdf_reader, "_render_pages_to_png", lambda data: [b"fake-png-bytes"])
    called = {}

    def fake_extract_from_images(document_bytes, images, role, filename):
        called["used_vision_path"] = True
        return DocumentExtraction(role=role, filename=filename, doc_kind="BL", readable=True)

    monkeypatch.setattr(pdf_reader, "extract_from_images", fake_extract_from_images)
    result = pdf_reader.read_pdf(b"fake-pdf-bytes", "BL", "att.pdf")

    assert called.get("used_vision_path") is True
    assert result.doc_kind == "BL"


def test_pdf_reader_escalates_when_text_is_truncated(monkeypatch):
    import readers.pdf as pdf_reader

    monkeypatch.setattr(pdf_reader, "_extract_text_per_page", lambda data: (["Shipper: TEST CO\n" * 5], True))

    result = pdf_reader.read_pdf(b"fake-pdf-bytes", "SI", "att.pdf")

    assert result.readable is False
    assert result.doc_kind == "UNKNOWN"


def test_pdf_reader_unreadable_when_no_text_and_no_images(monkeypatch):
    import readers.pdf as pdf_reader

    monkeypatch.setattr(pdf_reader, "_extract_text_per_page", lambda data: ([""], False))
    monkeypatch.setattr(pdf_reader, "_render_pages_to_png", lambda data: [])

    result = pdf_reader.read_pdf(b"corrupt-bytes", "SI", "att.pdf")
    assert result.readable is False
    assert result.doc_kind == "UNKNOWN"


def test_pdf_reader_uses_vision_when_only_one_page_has_footer_text(monkeypatch):
    # Repro: a scanned PDF where page 1 has a 40+ char footer/stamp and the
    # rest of the pages have no text at all. Routing on *total* text length
    # would wrongly use the text path and Gemini would only ever see the
    # footer, missing every real field.
    import readers.pdf as pdf_reader

    monkeypatch.setattr(
        pdf_reader,
        "_extract_text_per_page",
        lambda data: (["Page 1 of 3 - Document ID 0123456789ABCDEF", "", ""], False),
    )
    monkeypatch.setattr(pdf_reader, "_render_pages_to_png", lambda data: [b"p1", b"p2", b"p3"])
    called = {}

    def fake_extract_from_images(document_bytes, images, role, filename):
        called["used_vision_path"] = True
        return DocumentExtraction(role=role, filename=filename, doc_kind="BL", readable=True)

    monkeypatch.setattr(pdf_reader, "extract_from_images", fake_extract_from_images)
    result = pdf_reader.read_pdf(b"fake-pdf-bytes", "BL", "att.pdf")

    assert called.get("used_vision_path") is True
    assert result.doc_kind == "BL"


def test_pdf_reader_uses_text_path_only_when_every_page_has_real_text(monkeypatch):
    import readers.pdf as pdf_reader

    monkeypatch.setattr(
        pdf_reader,
        "_extract_text_per_page",
        lambda data: (["Shipper: TEST CO\n" * 5, "Consignee: TEST CO\n" * 5], False),
    )
    called = {}

    def fake_extract_from_text(document_bytes, document_text, role, filename):
        called["used_text_path"] = True
        return DocumentExtraction(role=role, filename=filename, doc_kind="SI", readable=True)

    monkeypatch.setattr(pdf_reader, "extract_from_text", fake_extract_from_text)
    result = pdf_reader.read_pdf(b"fake-pdf-bytes", "SI", "att.pdf")

    assert called.get("used_text_path") is True
    assert result.doc_kind == "SI"
